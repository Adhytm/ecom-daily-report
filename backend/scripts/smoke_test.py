"""全链路冒烟测试（规划 T12 / T14 —— 最终交付的唯一硬性门槛）。

从空库开始执行：
初始化 → 生成模拟数据 → 上传 4 平台文件 → 自动识别 → commit →
SKU 映射建议采纳 → 生成日报 → 下载 Excel → 获取摘要 →
断言 manifest.json 关键数字（±0.5%）。

用法：``python -m backend.scripts.smoke_test``
"""

from __future__ import annotations

import io
import json
import os
import tempfile
from decimal import Decimal
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[2]

# 允许的端到端数字误差（规划 6.8：±0.5%）
TOLERANCE = Decimal("0.005")


def _assert_close(actual: float, expected: Decimal, label: str) -> None:
    exp = Decimal(str(expected))
    if exp == 0:
        assert abs(Decimal(str(actual))) <= Decimal("0.01"), (
            f"{label}: 期望 0，实际 {actual}"
        )
        return
    diff = abs(Decimal(str(actual)) - exp) / exp
    assert diff <= TOLERANCE, f"{label}: 期望 {exp}，实际 {actual}，误差 {diff * 100:.3f}%"


def main() -> int:
    # ---- 0) 空库 + 临时目录（必须在导入 backend.app 之前设置环境变量） ----
    tmp = Path(tempfile.mkdtemp(prefix="ecom_smoke_"))
    os.environ["ECOM_DB_PATH"] = str(tmp / "app.db")
    os.environ["EXPORT_DIR"] = str(tmp / "exports")
    os.environ["SUMMARY_USE_WAN"] = "true"

    from fastapi.testclient import TestClient

    from backend.app.main import app

    with TestClient(app) as client:
        try:
            _run_pipeline(client, tmp)
            print("\n===== 冒烟测试全部通过 =====")
        finally:
            # 产物（SQLite 库 + 导出的 Excel）故意保留，事后排查离不开它；
            # 但必须告知位置，否则每跑一次就往 %TEMP% 静默堆一个 3MB+ 的库。
            print(f"排查产物已保留：{tmp}")
    return 0


def _run_pipeline(client: TestClient, tmp: Path) -> None:
    # ---- 1) 健康检查 ----
    resp = client.get("/api/health")
    assert resp.status_code == 200 and resp.json()["ok"] is True
    assert resp.json()["data"]["db"] == "ok"
    assert resp.json()["data"]["builtin_adapters"] == 4
    print("[1/9] 健康检查通过（4 个内置适配器）")

    # ---- 2) 生成模拟数据 ----
    resp = client.post("/api/mock-data/generate", json={"days": 35, "skus": 40})
    assert resp.status_code == 200, resp.text
    mock_out = Path(resp.json()["data"]["outdir"])
    manifest = json.loads((mock_out / "manifest.json").read_text(encoding="utf-8"))
    report_date = manifest["report_date"]
    print(f"[2/9] 模拟数据生成：4 平台文件 + manifest（报告日 {report_date}）")

    # ---- 3) 建内部 SKU 目录 ----
    for sku in manifest["internal_skus"]:
        resp = client.post("/api/skus", json={
            "sku_code": sku["sku_code"], "name": sku["name"],
            "category": sku["category"], "brand": sku.get("brand"),
        })
        assert resp.status_code == 200, resp.text
    print(f"[3/9] 内部 SKU 目录：{len(manifest['internal_skus'])} 个")

    # ---- 4) 上传 4 平台文件（自动识别平台） ----
    upload_ids: dict[str, int] = {}
    for platform_key, fname in manifest["files"].items():
        raw = (mock_out / fname).read_bytes()
        resp = client.post(
            "/api/uploads",
            files={"files": (fname, raw, "application/octet-stream")},
        )
        assert resp.status_code == 200, resp.text
        item = resp.json()["data"]["files"][0]
        assert item["detected_platform"] == platform_key, (
            f"{fname} 被识别为 {item['detected_platform']}，应为 {platform_key}"
        )
        assert item["stats"]["total"] > 0
        upload_ids[platform_key] = item["upload_id"]
    print("[4/9] 4 个平台文件上传且自动识别全部正确")

    # ---- 5) 提交入库（触发 SKU 映射） ----
    quality = {"filtered": 0, "duplicated": 0, "failed": 0}
    for platform_key, upload_id in upload_ids.items():
        resp = client.post(f"/api/uploads/{upload_id}/commit")
        assert resp.status_code == 200, resp.text
        data = resp.json()["data"]
        assert data["unmapped_count"] > 0  # 独占 SKU 保持未映射

        resp = client.get(f"/api/uploads/{upload_id}/errors?size=200")
        for e in resp.json()["data"]["items"]:
            if e["error_type"] == "filtered_out":
                quality["filtered"] += 1
            elif e["error_type"] == "duplicate_row":
                quality["duplicated"] += 1
            elif e["error_type"] in ("number_parse_failed", "date_parse_failed"):
                quality["failed"] += 1
    exp_q = manifest["expected"]["quality"]
    assert quality["filtered"] >= exp_q["filtered_min"], quality
    assert quality["duplicated"] >= exp_q["duplicated_min"], quality
    assert quality["failed"] >= exp_q["failed_min"], quality
    print(f"[5/9] 提交入库完成，行级错误统计：{quality}")

    # ---- 6) 模糊匹配建议 + 采纳高置信候选 ----
    resp = client.post("/api/sku-mappings/suggest", json={})
    items = resp.json()["data"]["items"]
    adopted = 0
    for item in items:
        for cand in item["candidates"]:
            if cand["score"] >= 90:
                resp = client.put(
                    f"/api/sku-mappings/{item['mapping_id']}",
                    json={"sku_id": cand["sku_id"]},
                )
                assert resp.status_code == 200, resp.text
                adopted += 1
                break
    assert adopted >= len(manifest["internal_skus"]), (
        f"采纳数 {adopted} 少于内部 SKU 数 {len(manifest['internal_skus'])}"
    )
    print(f"[6/9] 映射建议采纳：{adopted} 条（得分 >= 90）")

    # ---- 7) 看板聚合 + manifest 数字断言（±0.5%） ----
    resp = client.get(f"/api/reports/daily?date={report_date}")
    assert resp.status_code == 200, resp.text
    report = resp.json()["data"]

    # 报告日单日：overview（只含当日有数据的平台）vs manifest 报告日分平台预期
    exp_platforms = manifest["expected"]["report_day_by_platform"]
    by_platform = {r["key"]: r for r in report["by_platform"]}
    day_gmv = Decimal("0")
    day_refund = Decimal("0")
    for platform_key, bucket in exp_platforms.items():
        if platform_key == "jd":
            continue  # 京东报告日故意缺失
        assert platform_key in by_platform, f"{platform_key} 未出现在分平台"
        _assert_close(by_platform[platform_key]["metrics"]["gmv"]["value"],
                      bucket["gmv"], f"{platform_key} GMV（报告日）")
        day_gmv += Decimal(str(bucket["gmv"]))
        day_refund += Decimal(str(bucket["refund_amount"]))
    _assert_close(report["overview"]["gmv"]["value"], day_gmv, "总 GMV（报告日）")
    _assert_close(report["overview"]["refund_amount"]["value"], day_refund,
                  "总退款额（报告日）")
    _assert_close(report["overview"]["net_amount"]["value"], day_gmv - day_refund,
                  "总实际成交（报告日）")

    # 全周期（近 30 天趋势序列求和）vs manifest 按日 GMV
    trend = report["trend"]
    trend_gmv_sum = Decimal("0")
    for d, v in zip(trend["dates"], trend["series"]["gmv"]):
        if v is not None:
            trend_gmv_sum += Decimal(str(v))
    manifest_day_sum = Decimal("0")
    for d, v in manifest["expected"]["by_day_gmv"].items():
        if d in trend["dates"]:
            manifest_day_sum += Decimal(str(v))
    _assert_close(float(trend_gmv_sum), manifest_day_sum, "近 30 天 GMV 合计（趋势）")

    # data_completeness：京东最后一天缺失
    dc = report["data_completeness"]
    assert dc["missing_platforms"] == manifest["expected"]["missing_platforms_on_report_day"], dc

    # 京东含运费标注
    assert report["overview"].get("note") and "含运费" in report["overview"]["note"]

    # 未映射分组保留（【待映射】）
    assert len(report["by_sku"]["unmapped"]) > 0
    assert len(report["by_sku"]["mapped"]) == len(manifest["internal_skus"])
    print("[7/9] 看板聚合数字与 manifest 误差 <= 0.5%，缺失平台 / 含运费标注 / 待映射分组正确")

    # ---- 8) 异常预警：至少 4 条不同规则命中 ----
    rule_names = {a["rule_name"] for a in report["anomalies"]}
    assert len(rule_names) >= 4, f"仅命中 {rule_names}，需至少 4 条不同规则"
    print(f"[8/9] 异常预警命中 {len(rule_names)} 条不同规则：{sorted(rule_names)}")

    # ---- 9) 生成日报 → 下载 Excel → 摘要 ----
    resp = client.post("/api/reports/daily/generate", json={"date": report_date})
    assert resp.status_code == 200, resp.text
    gen = resp.json()["data"]
    assert len(gen["summary_text"]) == gen["char_count"]
    assert gen["char_count"] <= 800, f"摘要 {gen['char_count']} 字符超过 800"

    # 幂等：重复生成不新增
    resp2 = client.post("/api/reports/daily/generate", json={"date": report_date})
    assert resp2.json()["data"]["report_id"] == gen["report_id"]

    resp = client.get(f"/api/reports/daily/{report_date}/export")
    assert resp.status_code == 200, resp.text
    assert "filename*=UTF-8''" in resp.headers["content-disposition"]
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["日报总览", "分平台", "分店铺", "分类目", "分SKU",
                             "异常预警", "口径说明"]

    resp = client.get(f"/api/reports/daily/{report_date}/summary")
    assert resp.status_code == 200
    summary = resp.json()["data"]["summary_text"]
    assert "◆ 整体" in summary and "◆ 分平台" in summary
    print(f"[9/9] 日报生成（Excel 7 Sheet 可下载）+ 摘要 {gen['char_count']} 字符")

    # 历史日报
    resp = client.get("/api/reports")
    assert resp.json()["data"]["total"] == 1


if __name__ == "__main__":
    raise SystemExit(main())
