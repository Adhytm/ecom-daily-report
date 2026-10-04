"""Excel 日报导出（规划 6.6）。

固定 7 个 Sheet：日报总览 / 分平台 / 分店铺 / 分类目 / 分SKU /
异常预警 / 口径说明。样式：主题色 1F4E79 表头、隔行 F2F7FB、
冻结 B2、金额 #,##0.00、百分比 0.0%、变化率带 +/- 着色、
条件格式高亮环比超阈、隐藏网格线。
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---- 样式常量（规划 6.6） ----
HEADER_BG = "1F4E79"
BAND_BG = "F2F7FB"
WARN_BG = "FFF2CC"       # 【待映射】块底色
COND_BG = "FFC7CE"       # 条件格式：环比超阈
CRITICAL_BG = "C00000"   # critical 行红底白字
WARNING_BG = "FFEB9C"    # warning 行黄底
POS_COLOR = "006100"
NEG_COLOR = "9C0006"

FMT_MONEY = "#,##0.00"
FMT_INT = "#,##0"
FMT_PCT = "0.0%"
FMT_PCT_SIGNED = "+0.0%;-0.0%;0.0%"
FMT_PP_SIGNED = '+0.0"pp";-0.0"pp";0.0"pp"'

_THIN = Side(style="thin", color="D9D9D9")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _header_style(cell) -> None:
    cell.fill = PatternFill("solid", fgColor=HEADER_BG)
    cell.font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
    cell.alignment = Alignment(horizontal="center", vertical="center")


def _data_style(cell, banded: bool = False) -> None:
    cell.font = Font(name="微软雅黑", size=10)
    cell.border = _BORDER
    if banded:
        cell.fill = PatternFill("solid", fgColor=BAND_BG)


def _delta_style(cell, value: Any) -> None:
    """变化率单元格：正数绿色带 +，负数红色。"""
    if value is None:
        cell.value = "—"
        cell.alignment = Alignment(horizontal="center")
        return
    color = POS_COLOR if value > 0 else (NEG_COLOR if value < 0 else "000000")
    cell.font = Font(name="微软雅黑", size=10, color=color, bold=abs(value) >= 0.3)
    cell.alignment = Alignment(horizontal="right")


def _auto_width(ws, max_width: int = 40, min_width: int = 10) -> None:
    """按内容自适应列宽（CJK 计 2），上限 40 下限 10。"""
    for col_idx in range(1, ws.max_column + 1):
        width = min_width
        for row_idx in range(1, min(ws.max_row, 200) + 1):
            v = ws.cell(row=row_idx, column=col_idx).value
            if v is None:
                continue
            text = str(v)
            length = sum(2 if ord(ch) > 127 else 1 for ch in text[:60])
            width = max(width, min(max_width, length + 4))
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def _write_table(ws, start_row: int, headers: list[str], rows: list[list],
                 col_fmts: list[str | None], delta_cols: set[int] | None = None) -> int:
    """写一个带样式的表格，返回下一可用行号。delta_cols 为变化率列下标。"""
    delta_cols = delta_cols or set()
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=c, value=h)
        _header_style(cell)
        cell.border = _BORDER
    ws.row_dimensions[start_row].height = 24

    for r, row_vals in enumerate(rows, start=start_row + 1):
        banded = (r - start_row) % 2 == 0
        for c, v in enumerate(row_vals, start=1):
            cell = ws.cell(row=r, column=c)
            fmt = col_fmts[c - 1] if c - 1 < len(col_fmts) else None
            if c - 1 in delta_cols:
                _delta_style(cell, v)
                if v is not None:
                    cell.value = v
                    cell.number_format = fmt or FMT_PCT_SIGNED
            else:
                cell.value = v if v is not None else "—"
                _data_style(cell, banded)
                if fmt and v is not None:
                    cell.number_format = fmt
                if isinstance(v, str):
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                elif v is not None:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
    return start_row + 1 + len(rows)


# ---------------------------------------------------------------------------
# 各 Sheet 生成
# ---------------------------------------------------------------------------


def _sheet_overview(wb: Workbook, report: dict, company: str) -> None:
    ws = wb.active
    ws.title = "日报总览"
    ws.sheet_view.showGridLines = False

    ov = report["overview"]
    raw = report.get("overview_raw", {})
    dc = report["data_completeness"]

    ws.merge_cells("A1:F1")
    title = ws["A1"]
    title.value = f"{company} 电商销售日报 {report['report_date']} {report['weekday']}"
    title.font = Font(name="微软雅黑", size=14, bold=True, color=HEADER_BG)
    title.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    # (显示名, 指标键, 数值格式, 是否比率行, 环比格式)
    metric_rows_spec = [
        ("GMV", "gmv", FMT_MONEY, False, FMT_PCT_SIGNED),
        ("实际成交", "net_amount", FMT_MONEY, False, FMT_PCT_SIGNED),
        ("退款金额", "refund_amount", FMT_MONEY, False, FMT_PCT_SIGNED),
        ("退款率", "refund_rate", FMT_PCT, True, FMT_PP_SIGNED),
        ("支付件数", "paid_qty", FMT_INT, False, FMT_PCT_SIGNED),
        ("买家数", "buyers", FMT_INT, False, FMT_PCT_SIGNED),
        ("访客数", "visitors", FMT_INT, False, FMT_PCT_SIGNED),
        ("转化率", "conversion_rate", FMT_PCT, True, FMT_PP_SIGNED),
        ("客单价", "avg_order_value", FMT_MONEY, False, FMT_PCT_SIGNED),
        ("推广费", "ad_cost", FMT_MONEY, False, FMT_PCT_SIGNED),
        ("推广ROI", "roi", FMT_MONEY, False, FMT_PCT_SIGNED),
        ("毛利", "gross_profit", FMT_MONEY, False, FMT_PCT_SIGNED),
        ("毛利率", "gross_margin", FMT_PCT, True, FMT_PP_SIGNED),
        ("经营利润", "operating_profit", FMT_MONEY, False, FMT_PCT_SIGNED),
    ]

    headers = ["指标", "今日", "昨日", "环比", "上周同日", "周同比"]
    start = 3
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=start, column=c, value=h)
        _header_style(cell)
        cell.border = _BORDER
    ws.row_dimensions[start].height = 24

    for i, (label, key, fmt, is_ratio, delta_fmt) in enumerate(metric_rows_spec,
                                                               start=start + 1):
        banded = (i - start) % 2 == 0
        block = ov.get(key, {})
        r = raw.get(key, {})
        cells = [
            (label, None),
            (block.get("value"), fmt),
            (r.get("dod_value"), fmt),
            (block.get("dod_pp") if is_ratio else block.get("dod"), delta_fmt),
            (r.get("wow_value"), fmt),
            (block.get("wow_pp") if is_ratio else block.get("wow"), delta_fmt),
        ]
        for c, (v, cell_fmt) in enumerate(cells, start=1):
            cell = ws.cell(row=i, column=c)
            if c in (4, 6):
                _delta_style(cell, v)
                if v is not None:
                    cell.number_format = cell_fmt
            else:
                cell.value = v if v is not None else "—"
                _data_style(cell, banded)
                if v is not None and cell_fmt:
                    cell.number_format = cell_fmt
                if c > 1:
                    cell.alignment = Alignment(horizontal="right", vertical="center")

    note_row = start + len(metric_rows_spec) + 2
    if dc.get("warning"):
        ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=6)
        cell = ws.cell(row=note_row, column=1, value=dc["warning"])
        cell.fill = PatternFill("solid", fgColor="FFC7CE")
        cell.font = Font(name="微软雅黑", size=10, color="9C0006", bold=True)
        note_row += 1
    if ov.get("note"):
        ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=6)
        cell = ws.cell(row=note_row, column=1, value=f"口径备注：{ov['note']}")
        cell.font = Font(name="微软雅黑", size=10, color="833C00")
    _auto_width(ws)


def _dim_sheet(wb: Workbook, title: str, headers: list[str], rows: list[list],
               fmts: list[str | None], delta_cols: set[int], dod_col: int,
               first_cols: int = 1) -> None:
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    next_row = _write_table(ws, 1, headers, rows, fmts, delta_cols)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(next_row - 1, 2)}"
    # 条件格式：GMV 环比列绝对值 > 30% 高亮
    col_letter = get_column_letter(dod_col)
    rng = f"{col_letter}2:{col_letter}{max(next_row - 1, 2)}"
    ws.conditional_formatting.add(
        rng,
        FormulaRule(formula=[f"AND(ISNUMBER({col_letter}2),ABS({col_letter}2)>0.3)"],
                    fill=PatternFill("solid", fgColor=COND_BG)),
    )
    _auto_width(ws)


def build_platform_rows(report: dict) -> tuple[list[str], list[list], list[str | None], set[int], int]:
    headers = ["平台", "GMV", "占比", "实际成交", "退款额", "退款率", "件数", "买家数",
               "访客", "转化率", "客单价", "推广费", "ROI", "GMV环比", "GMV周同比"]
    rows = []
    for r in report["by_platform"]:
        m = r["metrics"]
        rows.append([
            r["name"],
            m["gmv"]["value"], r.get("share"), m["net_amount"]["value"],
            m["refund_amount"]["value"], m["refund_rate"]["value"],
            m["paid_qty"]["value"], m["buyers"]["value"], m["visitors"]["value"],
            m["conversion_rate"]["value"], m["avg_order_value"]["value"],
            m["ad_cost"]["value"], m["roi"]["value"],
            m["gmv"]["dod"], m["gmv"]["wow"],
        ])
    fmts = [None, FMT_MONEY, FMT_PCT, FMT_MONEY, FMT_MONEY, FMT_PCT, FMT_INT, FMT_INT,
            FMT_INT, FMT_PCT, FMT_MONEY, FMT_MONEY, FMT_MONEY, FMT_PCT_SIGNED, FMT_PCT_SIGNED]
    return headers, rows, fmts, {13, 14}, 14


def build_shop_rows(report: dict):
    headers = ["平台", "店铺", "GMV", "占比", "实际成交", "退款额", "退款率", "件数",
               "买家数", "访客", "转化率", "客单价", "推广费", "ROI", "GMV环比", "GMV周同比"]
    rows = []
    for r in report["by_shop"]:
        m = r["metrics"]
        parts = r["name"].split(" · ", 1)
        platform = parts[0] if len(parts) > 1 else ""
        shop = parts[1] if len(parts) > 1 else r["name"]
        rows.append([
            platform, shop,
            m["gmv"]["value"], r.get("share"), m["net_amount"]["value"],
            m["refund_amount"]["value"], m["refund_rate"]["value"],
            m["paid_qty"]["value"], m["buyers"]["value"], m["visitors"]["value"],
            m["conversion_rate"]["value"], m["avg_order_value"]["value"],
            m["ad_cost"]["value"], m["roi"]["value"],
            m["gmv"]["dod"], m["gmv"]["wow"],
        ])
    fmts = [None, None, FMT_MONEY, FMT_PCT, FMT_MONEY, FMT_MONEY, FMT_PCT, FMT_INT,
            FMT_INT, FMT_INT, FMT_PCT, FMT_MONEY, FMT_MONEY, FMT_MONEY,
            FMT_PCT_SIGNED, FMT_PCT_SIGNED]
    return headers, rows, fmts, {14, 15}, 15


def build_category_rows(report: dict):
    headers = ["类目", "GMV", "占比", "实际成交", "件数", "件单价", "GMV环比", "GMV周同比"]
    rows = []
    for r in report["by_category"]:
        m = r["metrics"]
        rows.append([
            r["name"], m["gmv"]["value"], r.get("share"), m["net_amount"]["value"],
            m["paid_qty"]["value"], m.get("unit_price", {}).get("value") if isinstance(
                m.get("unit_price"), dict) else None,
            m["gmv"]["dod"], m["gmv"]["wow"],
        ])
    fmts = [None, FMT_MONEY, FMT_PCT, FMT_MONEY, FMT_INT, FMT_MONEY,
            FMT_PCT_SIGNED, FMT_PCT_SIGNED]
    return headers, rows, fmts, {6, 7}, 7


def build_sku_rows(report: dict):
    """分SKU：已映射在前（按实际成交降序），【待映射】分组在末尾。"""
    headers = ["SKU", "商品名", "类目", "平台", "GMV", "实际成交", "件数", "退款率",
               "毛利", "GMV环比"]
    rows: list[list] = []
    unmapped_start: int | None = None
    for r in report["by_sku"]["mapped"]:
        m = r["metrics"]
        rows.append([
            r["sku_code"], r["name"], r.get("category") or "—", r.get("platforms") or "—",
            m["gmv"]["value"], m["net_amount"]["value"], m["paid_qty"]["value"],
            m["refund_rate"]["value"], m["gross_profit"]["value"], m["gmv"]["dod"],
        ])
    unmapped_start = len(rows) + 2  # +表头行
    for r in report["by_sku"]["unmapped"]:
        m = r["metrics"]
        try:
            from backend.app.core.aggregator import _platform_display

            platform = _platform_display(r["platform_key"])
        except Exception:  # noqa: BLE001
            platform = r["platform_key"]
        rows.append([
            "【待映射】", r.get("product_name") or r["platform_product_code"],
            "—", platform,
            m["gmv"]["value"], m["net_amount"]["value"], m["paid_qty"]["value"],
            m["refund_rate"]["value"], m["gross_profit"]["value"], m["gmv"]["dod"],
        ])
    fmts = [None, None, None, None, FMT_MONEY, FMT_MONEY, FMT_INT, FMT_PCT,
            FMT_MONEY, FMT_PCT_SIGNED]
    return headers, rows, fmts, {9}, 10, unmapped_start


def _sheet_anomaly(wb: Workbook, report: dict) -> None:
    ws = wb.create_sheet("异常预警")
    ws.sheet_view.showGridLines = False
    headers = ["级别", "范围", "对象", "指标", "当前值", "阈值", "说明"]
    anomalies = report.get("anomalies", [])
    if not anomalies:
        rows = [["—", "—", "—", "—", "—", "—", "本日无异常"]]
    else:
        rows = [
            [
                "critical" if a["severity"] == "critical" else "warning",
                a["scope"], a["subject"], a["metric"],
                a["value"] if a["value"] is not None else "—",
                a["threshold"], a["message"],
            ]
            for a in anomalies
        ]
    next_row = _write_table(ws, 1, headers, rows, [None] * 7)
    for idx, a in enumerate(anomalies, start=2):
        fill = CRITICAL_BG if a["severity"] == "critical" else WARNING_BG
        font_color = "FFFFFF" if a["severity"] == "critical" else "7F6000"
        for c in range(1, 8):
            cell = ws.cell(row=idx, column=c)
            cell.fill = PatternFill("solid", fgColor=fill)
            if a["severity"] == "critical":
                cell.font = Font(name="微软雅黑", size=10, color=font_color, bold=True)
    _auto_width(ws)


def _sheet_semantics(wb: Workbook, report: dict, uploads_info: list[dict]) -> None:
    ws = wb.create_sheet("口径说明")
    ws.sheet_view.showGridLines = False

    row = 1
    def _title(text: str) -> None:
        nonlocal row
        cell = ws.cell(row=row, column=1, value=text)
        cell.font = Font(name="微软雅黑", size=12, bold=True, color=HEADER_BG)
        row += 1

    def _line(text: str) -> None:
        nonlocal row
        ws.cell(row=row, column=1, value=text).font = Font(name="微软雅黑", size=10)
        row += 1

    dc = report["data_completeness"]
    _title("一、数据来源")
    if uploads_info:
        for u in uploads_info:
            _line(
                f"· {u['filename']}（平台：{u['platform_key'] or '未知'}，"
                f"上传时间：{u['created_at']}，总行数：{u['row_count_total']}，"
                f"有效行数：{u['row_count_valid']}）"
            )
    else:
        _line("· 本地生成数据（无上传记录）")
    if dc.get("warning"):
        _line(f"· 缺失数据：{dc['warning']}")

    row += 1
    _title("二、指标定义")
    for line in [
        "· GMV：支付金额（口径见下表平台差异）",
        "· 实际成交额 = GMV - 退款金额（按各平台 net_amount_formula 计算）",
        "· 退款率 = 退款金额 / GMV",
        "· 客单价 = 实际成交额 / 支付买家数",
        "· 件单价 = 实际成交额 / 支付件数",
        "· 转化率 = 支付买家数 / 访客数",
        "· UV 价值 = 实际成交额 / 访客数",
        "· 推广 ROI = 实际成交额 / 推广花费",
        "· 推广费比 = 推广花费 / 实际成交额",
        "· 毛利 = 实际成交额 - 货物成本（SKU成本价×件数）- 平台扣点 - 单件履约成本",
        "· 毛利率 = 毛利 / 实际成交额",
        "· 经营利润 = 毛利 - 推广费（未含人力、仓储等固定开支）",
        "· 平台扣点 / 单件履约成本在各平台适配器 YAML 的 metric_semantics 中配置",
        "· 环比 = (今日 - 昨日) / 昨日",
        "· 周同比 = (今日 - 上周同日) / 上周同日（电商报表统一采用周同比，消除工作日效应）",
        "· 比率类指标（退款率/转化率）的环比用百分点差值（pp）表示",
    ]:
        _line(line)

    row += 1
    _title("三、平台口径差异")
    diff_headers = ["平台", "GMV 口径", "含运费", "含税", "退款处理", "实际成交公式", "数据粒度"]
    for c, h in enumerate(diff_headers, start=1):
        cell = ws.cell(row=row, column=c, value=h)
        _header_style(cell)
    ws.row_dimensions[row].height = 24
    header_row = row
    row += 1

    from backend.app.core.adapter_registry import list_builtin_keys, load_builtin

    basis_text = {"paid": "支付", "ordered": "下单", "shipped": "发货"}
    refund_text = {"separate": "单独字段扣除", "deducted": "已扣除", "none": "无退款数据"}
    for key in list_builtin_keys():
        try:
            spec = load_builtin(key).spec
        except Exception:  # noqa: BLE001
            continue
        sem = spec.metric_semantics
        vals = [
            spec.display_name,
            basis_text[sem.gmv_basis],
            "是" if sem.gmv_includes_shipping else "否",
            "是" if sem.gmv_includes_tax else "否",
            refund_text[sem.refund_handling],
            sem.net_amount_formula,
            sem.grain,
        ]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=row, column=c, value=v)
            _data_style(cell, (row - header_row) % 2 == 0)
        row += 1

    row += 1
    _line("注意：京东 GMV 含运费，与其他平台不可直接相加比较；")
    _line("全平台合计 GMV 照常相加，但结果旁标注 *京东GMV含运费。")
    _auto_width(ws)


# ---------------------------------------------------------------------------
# 主入口
# ---------------------------------------------------------------------------


def generate_excel_report(report: dict, out_dir: Path | str,
                          uploads_info: list[dict] | None = None,
                          company: str = "XX") -> str:
    """生成日报 Excel，返回文件绝对路径。"""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    uploads_info = uploads_info or []

    wb = Workbook()
    _sheet_overview(wb, report, company)

    headers, rows, fmts, deltas, dod_col = build_platform_rows(report)
    _dim_sheet(wb, "分平台", headers, rows, fmts, deltas, dod_col)

    headers, rows, fmts, deltas, dod_col = build_shop_rows(report)
    _dim_sheet(wb, "分店铺", headers, rows, fmts, deltas, dod_col)

    headers, rows, fmts, deltas, dod_col = build_category_rows(report)
    _dim_sheet(wb, "分类目", headers, rows, fmts, deltas, dod_col)

    headers, rows, fmts, deltas, dod_col, unmapped_start = build_sku_rows(report)
    ws = wb.create_sheet("分SKU")
    ws.sheet_view.showGridLines = False
    next_row = _write_table(ws, 1, headers, rows, fmts, deltas)
    if unmapped_start is not None:
        for r in range(unmapped_start, next_row):
            for c in range(1, len(headers) + 1):
                ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor=WARN_BG)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(next_row - 1, 2)}"
    col_letter = get_column_letter(dod_col)
    ws.conditional_formatting.add(
        f"{col_letter}2:{col_letter}{max(next_row - 1, 2)}",
        FormulaRule(formula=[f"AND(ISNUMBER({col_letter}2),ABS({col_letter}2)>0.3)"],
                    fill=PatternFill("solid", fgColor=COND_BG)),
    )
    # 商品名列固定宽度 45
    ws.column_dimensions["B"].width = 45
    _auto_width(ws)
    ws.column_dimensions["B"].width = 45

    _sheet_anomaly(wb, report)
    _sheet_semantics(wb, report, uploads_info)

    stamp = datetime.now()
    path = out_dir / f"daily_report_{stamp.strftime('%Y%m%d')}_{stamp.strftime('%H%M%S')}.xlsx"
    wb.save(path)
    return str(path)
