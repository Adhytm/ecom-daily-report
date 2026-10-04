"""文件解析层（规划 6.1）。

把上传文件的字节流变成 pandas.DataFrame（**原始表头、未映射**）。
本模块只知道"物理层"：编码、表头定位、表头清洗、空行/尾行丢弃，
完全不感知任何平台业务字段；所有类型转换是 normalizer 的职责。
"""

from __future__ import annotations

import csv
import io
from typing import Any

import pandas as pd
from charset_normalizer import from_bytes
from openpyxl import load_workbook

from backend.app.core import EncodingError, FileParseError
from backend.app.schemas.adapter import FileMatching
from backend.app.schemas.upload import ParseMeta

# 不可见字符与换行符（表头清洗用）
_INVISIBLE_CHARS = ("﻿", "​", "‌", "‍", "⁠", "﻿")
_NEWLINES = ("\r\n", "\r", "\n")

AUTO_SCAN_LIMIT = 200  # auto 模式最多扫描前 200 行

# 全局支持的文件扩展名（单一事实来源；API / 自动接入服务统一引用）
SUPPORTED_EXTENSIONS: tuple[str, ...] = (".csv", ".xlsx", ".xlsm", ".xls", ".tsv", ".txt")

_XLS_BIFF_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"  # OLE2 复合文档头（老式 .xls）
_XLSX_ZIP_MAGIC = b"PK\x03\x04"                       # ZIP 头（xlsx / xlsm）


def _is_biff(raw: bytes) -> bool:
    """老式 .xls（BIFF）是 OLE2 复合文档，不是 zip。"""
    return raw[:8] == _XLS_BIFF_MAGIC


def _allowed_extensions(matching: FileMatching) -> tuple[str, ...]:
    """平台适配器显式声明的扩展名优先，否则用全局清单。"""
    declared = tuple(e.lower() for e in (matching.extensions or ()))
    return declared or SUPPORTED_EXTENSIONS


def _clean_header_cell(raw) -> str:
    """表头单元格清洗：strip、去换行、去不可见字符。"""
    text = "" if raw is None else str(raw)
    for ch in _NEWLINES:
        text = text.replace(ch, "")
    for ch in _INVISIBLE_CHARS:
        text = text.replace(ch, "")
    return text.strip()


def _dedup_headers(headers: list[str], warnings: list[str]) -> list[str]:
    """重复表头名自动加后缀 _2、_3，并记录警告。"""
    seen: dict[str, int] = {}
    result: list[str] = []
    for h in headers:
        count = seen.get(h, 0) + 1
        seen[h] = count
        result.append(h if count == 1 else f"{h}_{count}")
        if count == 2:
            warnings.append(f"检测到重复表头名“{h}”，已自动更名为“{h}_2”")
    return result


def _decode_csv(raw: bytes, filename: str, candidates: list[str]) -> tuple[str, str]:
    """按候选编码顺序尝试解码；charset_normalizer 兜底；全失败抛 EncodingError。"""
    for enc in candidates:
        try:
            text = raw.decode(enc)
            return text, enc
        except (UnicodeDecodeError, LookupError):
            continue

    # 兜底：charset_normalizer 自动探测
    best = from_bytes(raw).best()
    if best is not None:
        try:
            return str(best), best.encoding or "unknown"
        except (UnicodeDecodeError, ValueError):
            pass

    raise EncodingError(
        f"无法识别文件编码：{filename}",
        detail={"tried_encodings": candidates},
    )


def _sniff_delimiter(text: str) -> str:
    """从前若干非空行嗅探分隔符：制表符优先，其次逗号，兜底逗号。"""
    sample = [ln for ln in text.split("\n")[:20] if ln.strip()][:10]
    if not sample:
        return ","
    tab_score = sum(ln.count("\t") for ln in sample)
    comma_score = sum(ln.count(",") for ln in sample)
    return "\t" if tab_score > comma_score else ","


def _read_csv_rows(raw: bytes, filename: str, matching: FileMatching):
    text, encoding = _decode_csv(raw, filename, matching.encoding_candidates)
    # 剥离 BOM、统一换行符（规划 6.1 第 3 步）
    text = text.lstrip("﻿").replace("\r\n", "\n").replace("\r", "\n")
    delimiter = _sniff_delimiter(text)
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    return rows, encoding


def _pick_sheet(wb, matching: FileMatching, warnings: list[str]):
    """按 matching.sheet 选出工作表；auto 时挑"最像数据表"的那个。

    打分规则：auto_header_hints 命中数 > 0 的候选里取命中最多者；
    **全部落空时回退首个工作表**（与 sheet=None 一致），只在 warnings 里
    说明结果可能不对——多 Sheet 文件不得因为"没有一页命中提示词"而整份
    报废（改造前默认取首个工作表，这是既有兼容行为）。
    """
    sheets = wb.worksheets
    if not sheets:
        raise FileParseError("Excel 文件中没有任何工作表")
    sheet = matching.sheet
    if isinstance(sheet, str) and sheet != "auto":
        if sheet not in wb.sheetnames:
            raise FileParseError(
                f"工作表“{sheet}”不存在",
                detail={"available_sheets": wb.sheetnames},
            )
        return wb[sheet]
    if isinstance(sheet, list) and sheet:
        for name in sheet:
            if name in wb.sheetnames:
                return wb[name]
        raise FileParseError(
            f"指定的工作表均不存在：{sheet}",
            detail={"available_sheets": wb.sheetnames},
        )
    if sheet is None:
        return sheets[0]

    # auto：按表头命中数挑最像数据表的工作表
    hints = set(matching.auto_header_hints)
    best, best_score = None, -1
    for ws in sheets:
        score = 0
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= AUTO_SCAN_LIMIT:
                break
            score = max(
                score,
                sum(1 for c in row if _clean_header_cell(c) in hints),
            )
        if score > best_score:
            best, best_score = ws, score
    if best is not None and best_score > 0:
        return best

    fallback = sheets[0]
    warnings.append(
        f"没有任何工作表的表头命中 auto_header_hints {sorted(hints)}，"
        f"已回退首个工作表“{fallback.title}”（可用工作表："
        f"{'、'.join(wb.sheetnames)}）；如结果不对，请在平台适配器的 "
        f"file_matching.sheet 指定工作表"
    )
    return fallback


def _read_xlsx_rows(
    raw: bytes, matching: FileMatching, warnings: list[str]
) -> tuple[list[list], str, list[str]]:
    """openpyxl 只读模式读取工作表全部行（值一律转字符串）。

    返回 (rows, 实际读取的工作表名, 文件内全部工作表名)。
    """
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 — 统一收敛为业务错误
        raise FileParseError(
            "无法读取该 Excel 文件：文件可能已损坏，或其实是老式 .xls 格式"
            "（请用 Excel/WPS 打开后另存为 .xlsx 再上传）",
            detail={"reason": type(e).__name__, "message": str(e)[:200]},
        ) from e
    try:
        ws = _pick_sheet(wb, matching, warnings)
        rows = [
            ["" if v is None else str(v) for v in row]
            for row in ws.iter_rows(values_only=True)
        ]
        return rows, ws.title, list(wb.sheetnames)
    finally:
        wb.close()


def _read_xls_rows(
    raw: bytes, matching: FileMatching, warnings: list[str]
) -> tuple[list[list], str, list[str]]:
    """xlrd 读取 .xls（BIFF）工作表全部行（值一律转字符串）。

    返回 (rows, 实际读取的工作表名, 文件内全部工作表名)。
    """
    import xlrd  # 延迟导入：仅 .xls 路径需要

    try:
        book = xlrd.open_workbook(file_contents=raw)
    except Exception as e:  # noqa: BLE001 — 统一收敛为业务错误
        raise FileParseError(
            "无法读取该 .xls 文件：文件可能已损坏",
            detail={"reason": type(e).__name__, "message": str(e)[:200]},
        ) from e
    try:
        names = book.sheet_names()
        if book.nsheets == 0:
            raise FileParseError("Excel 文件中没有任何工作表")
        sheet = matching.sheet
        target = None
        if isinstance(sheet, str) and sheet != "auto":
            if sheet not in names:
                raise FileParseError(
                    f"工作表“{sheet}”不存在",
                    detail={"available_sheets": names},
                )
            target = book.sheet_by_name(sheet)
        elif isinstance(sheet, list) and sheet:
            for name in sheet:
                if name in names:
                    target = book.sheet_by_name(name)
                    break
            if target is None:
                raise FileParseError(
                    f"指定的工作表均不存在：{sheet}",
                    detail={"available_sheets": names},
                )
        elif sheet is None:
            target = book.sheet_by_index(0)
        else:  # auto：按表头命中数挑最像数据表的工作表
            hints = set(matching.auto_header_hints)
            best, best_score = None, -1
            for name in names:
                sh = book.sheet_by_name(name)
                score = 0
                for r in range(min(sh.nrows, AUTO_SCAN_LIMIT)):
                    score = max(
                        score,
                        sum(1 for c in sh.row_values(r) if _clean_header_cell(c) in hints),
                    )
                if score > best_score:
                    best, best_score = sh, score
            if best is not None and best_score > 0:
                target = best
            else:
                # 与 _pick_sheet 同规则：全部落空回退首个工作表，不整份报废
                target = book.sheet_by_index(0)
                warnings.append(
                    f"没有任何工作表的表头命中 auto_header_hints {sorted(hints)}，"
                    f"已回退首个工作表“{target.name}”（可用工作表："
                    f"{'、'.join(names)}）；如结果不对，请在平台适配器的 "
                    f"file_matching.sheet 指定工作表"
                )
        rows = [
            ["" if v is None else str(v) for v in target.row_values(r)]
            for r in range(target.nrows)
        ]
        return rows, target.name, names
    finally:
        book.release_resources()


def _best_effort_header_index(rows: list[list], limit: int) -> int:
    """零命中兜底：猜一个"最像表头"的行下标（0-based）。

    规则：取首个「≥2 个非空单元格、且其后仍存在数据行」的行——
    这样能跳过只有标题的单列表头说明行，也能跳过纯空行。
    一个都没有时兜底首行；空文件返回 0（由调用方按空表处理）。
    """
    def _nonempty(row: list) -> int:
        return sum(1 for c in row if _clean_header_cell(c))

    scan = min(len(rows), limit)  # 表头候选行的扫描范围
    data_end = min(len(rows), limit + 1)  # 判断"其后有数据"时可多看一行
    for i in range(scan):
        if _nonempty(rows[i]) < 2:
            continue
        if any(_nonempty(rows[j]) >= 2 for j in range(i + 1, data_end)):
            return i
    return 0


def _locate_header_row(
    rows: list[list], matching: FileMatching, warnings: list[str]
) -> tuple[int, list[str]]:
    """返回 (表头行 1-based 序号, 清洗后的表头)。

    - header_row 为整数：直接采用（不足则补空列）
    - auto：扫描前 AUTO_SCAN_LIMIT 行，取与 auto_header_hints 命中数最多的行。
      **零命中不再抛错**：降级为 best-effort 猜表头行并写入 warnings，
      让调用方（平台识别 / 归一化）自己报出更准确的业务错误。
      改造前是"尽力而为"，强校验会让所有不含平台特征词的表头（自建导出、
      改名导出、本系统自产报表）在识别阶段就被拦死，属于兼容性回归。
    """
    if matching.header_row != "auto":
        idx = int(matching.header_row) - 1
        if idx < 0 or idx >= len(rows):
            raise FileParseError(
                f"指定的表头行 {matching.header_row} 超出文件行数（共 {len(rows)} 行）"
            )
        return int(matching.header_row), [
            _clean_header_cell(c) for c in rows[idx]
        ]

    best_idx, best_hits = -1, 0
    hints = set(matching.auto_header_hints)
    for i, row in enumerate(rows[:AUTO_SCAN_LIMIT]):
        cells = [_clean_header_cell(c) for c in row]
        hits = sum(1 for c in cells if c in hints)
        if hits > best_hits:
            best_idx, best_hits = i, hits
    if best_hits > 0:
        return best_idx + 1, [_clean_header_cell(c) for c in rows[best_idx]]

    # 零命中：降级为 best-effort，把"为什么猜错"的信息留在 warnings 里
    idx = _best_effort_header_index(rows, AUTO_SCAN_LIMIT)
    preview: list[Any] = [
        [_clean_header_cell(c) for c in row][:8] for row in rows[:5]
    ]
    warnings.append(
        f"前 {AUTO_SCAN_LIMIT} 行中没有任何一行命中 auto_header_hints "
        f"{sorted(hints)}，已按「首个含 ≥2 列且其后有数据行」的行推断表头为第 "
        f"{idx + 1} 行；如解析结果不对，请检查导出文件表头名称或适配器提示词。"
        f"文件前几行：{preview}"
    )
    return idx + 1, [_clean_header_cell(c) for c in rows[idx]]


def parse_file(raw: bytes, filename: str, matching: FileMatching) -> tuple[pd.DataFrame, ParseMeta]:
    """解析上传文件为 DataFrame（原始表头）+ ParseMeta。

    全部列保持 object/str 类型；NaN 不会出现（空单元格为空字符串）。
    格式路由按**内容（magic bytes）**而非后缀：xlsx 改名 .xls 也能正常解析。
    """
    lower = filename.lower()
    allowed = _allowed_extensions(matching)
    if not any(lower.endswith(ext) for ext in allowed):
        raise FileParseError(
            f"不支持的文件类型：{filename}",
            detail={"supported": list(allowed), "hint": "可在平台适配器中增补 extensions"},
        )

    warnings: list[str] = []
    sheet_name: str | None = None
    sheet_names: list[str] = []

    if lower.endswith((".xlsx", ".xlsm")) or (lower.endswith(".xls") and not _is_biff(raw)):
        # .xls 后缀但内容其实是 zip → 按 xlsx 读，避免用户改错后缀就报错
        rows, sheet_name, sheet_names = _read_xlsx_rows(raw, matching, warnings)
        encoding = "xlsx"
    elif _is_biff(raw):
        rows, sheet_name, sheet_names = _read_xls_rows(raw, matching, warnings)
        encoding = "xls"
    else:
        rows, encoding = _read_csv_rows(raw, filename, matching)

    header_row_index, headers = _locate_header_row(rows, matching, warnings)
    headers = _dedup_headers(headers, warnings)

    width = len(headers)
    data_rows = rows[header_row_index:]

    # 全空行丢弃
    dropped_empty = 0
    kept_rows: list[list[str]] = []
    for row in data_rows:
        # 行宽不齐：右侧补空串 / 截断到表头宽度
        cells = list(row[:width]) + [""] * (width - len(row))
        if all(str(c).strip() == "" for c in cells):
            dropped_empty += 1
            continue
        kept_rows.append(cells)

    # skip_footer：丢弃尾部 N 行
    if matching.skip_footer > 0:
        if matching.skip_footer >= len(kept_rows):
            warnings.append(
                f"skip_footer={matching.skip_footer} 不小于数据行数 {len(kept_rows)}，已全部丢弃"
            )
            kept_rows = []
        else:
            kept_rows = kept_rows[: len(kept_rows) - matching.skip_footer]
            warnings.append(f"已按配置丢弃文件尾部 {matching.skip_footer} 行")

    df = pd.DataFrame(kept_rows, columns=headers, dtype=object)

    meta = ParseMeta(
        encoding=encoding,
        header_row_index=header_row_index,
        detected_columns=list(headers),
        dropped_empty_rows=dropped_empty,
        warnings=warnings,
        sheet_name=sheet_name,
        sheet_names=sheet_names,
    )
    return df, meta
