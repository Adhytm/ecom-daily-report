"""Excel 日报导出测试（T10）：Sheet 结构、样式、条件格式、待映射标黄。"""

from __future__ import annotations

from datetime import date, timedelta
from decimal import Decimal

import openpyxl
import pytest

from backend.app.exporters.excel_report import generate_excel_report
from backend.app.models.entities import SalesFact, Sku, Upload

D0 = date(2026, 9, 6)
D1 = D0 - timedelta(days=1)
LONG_NAME = "超长商品名" + "很长的描述" * 20  # > 100 字符


@pytest.fixture()
def seeded_report(db_session):
    """构造含未映射、负增长、异常的已知数据集，返回聚合结果。"""
    up = Upload(filename="seed.csv", sha256="xlsx001", status="committed")
    db_session.add(up)
    sku_ok = Sku(sku_code="SKU1001", name="正常商品", category="美妆")
    sku_bad = Sku(sku_code="SKU1002", name=LONG_NAME, category="食品")
    db_session.add_all([sku_ok, sku_bad])
    db_session.flush()

    def fact(stat_date, platform, code, gmv, refund="0", sku_id=None, product=None):
        db_session.add(
            SalesFact(
                upload_id=up.id, stat_date=stat_date, platform_key=platform,
                shop_code="S1", shop_name="店铺", platform_product_code=code,
                product_name=product or f"商品{code}", category="美妆", sku_id=sku_id,
                order_qty=1, paid_qty=1, gmv=Decimal(str(gmv)),
                refund_amount=Decimal(str(refund)),
                net_amount=Decimal(str(gmv)) - Decimal(str(refund)),
                visitors=10, buyers=5, raw_row_json={},
            )
        )

    fact(D0, "taobao", "SKU1001", "1000", "100", sku_id=sku_ok.id)
    fact(D1, "taobao", "SKU1001", "2000", "100", sku_id=sku_ok.id)  # 环比 -50%
    fact(D0, "jd", "SKU1002", "800", "10", sku_id=sku_bad.id, product=LONG_NAME)
    fact(D0, "taobao", "UNMAPPED01", "300", "0")  # 未映射
    db_session.commit()

    from backend.app.core.aggregator import build_daily_report

    return build_daily_report(db_session, D0)


def test_seven_sheets_in_order(seeded_report, export_dir):
    path = generate_excel_report(seeded_report, export_dir, company="测试公司")
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["日报总览", "分平台", "分店铺", "分类目", "分SKU",
                             "异常预警", "口径说明"]


def test_header_style_and_freeze(seeded_report, export_dir):
    path = generate_excel_report(seeded_report, export_dir)
    wb = openpyxl.load_workbook(path)
    for name in ["分平台", "分SKU"]:
        ws = wb[name]
        header = ws.cell(row=1, column=1)
        assert header.fill.start_color.rgb.endswith("1F4E79")
        assert header.font.bold is True
        assert ws.freeze_panes == "B2"
        assert ws.auto_filter.ref is not None
        assert ws.sheet_view.showGridLines is False


def test_money_number_format(seeded_report, export_dir):
    path = generate_excel_report(seeded_report, export_dir)
    wb = openpyxl.load_workbook(path)
    ws = wb["分平台"]
    # B2 = GMV 列第一数据行
    cell = ws.cell(row=2, column=2)
    assert cell.number_format == "#,##0.00"
    # 退款率列 0.0%
    ws2 = wb["日报总览"]
    assert ws2.cell(row=4, column=2).value is not None


def test_negative_delta_red(seeded_report, export_dir):
    """环比 -50% 的 SKU 行：负数变化率字体色 9C0006。"""
    path = generate_excel_report(seeded_report, export_dir)
    wb = openpyxl.load_workbook(path)
    ws = wb["分SKU"]
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if cell.number_format == "+0.0%;-0.0%;0.0%" and isinstance(cell.value, (int, float)):
                if cell.value < 0:
                    assert cell.font.color.rgb.endswith("9C0006")
                    return
    pytest.fail("未找到负变化率单元格")


def test_unmapped_block_yellow(seeded_report, export_dir):
    """【待映射】分组整块底色 FFF2CC。"""
    path = generate_excel_report(seeded_report, export_dir)
    wb = openpyxl.load_workbook(path)
    ws = wb["分SKU"]
    found = False
    for row in ws.iter_rows(min_row=2):
        if row[0].value == "【待映射】":
            found = True
            for cell in row:
                assert cell.fill.start_color.rgb.endswith("FFF2CC")
            break
    assert found


def test_critical_row_red(seeded_report, export_dir):
    """critical 异常行红底白字。"""
    seeded_report["anomalies"] = [
        {"rule_name": "测试", "severity": "critical", "scope": "sku", "subject": "S",
         "metric": "gmv_dod", "value": -0.9, "threshold": -0.4, "message": "测试消息"}
    ]
    path = generate_excel_report(seeded_report, export_dir)
    wb = openpyxl.load_workbook(path)
    ws = wb["异常预警"]
    cell = ws.cell(row=2, column=1)
    assert cell.fill.start_color.rgb.endswith("C00000")
    assert cell.font.color.rgb.endswith("FFFFFF")


def test_no_anomaly_row(seeded_report, export_dir):
    seeded_report["anomalies"] = []
    path = generate_excel_report(seeded_report, export_dir)
    wb = openpyxl.load_workbook(path)
    ws = wb["异常预警"]
    values = [ws.cell(row=r, column=7).value for r in range(2, ws.max_row + 1)]
    assert "本日无异常" in values


def test_semantics_sheet_contains_platform_diff(seeded_report, export_dir):
    """口径说明 Sheet 含 4 个平台口径差异与京东含运费提示。"""
    path = generate_excel_report(seeded_report, export_dir)
    wb = openpyxl.load_workbook(path)
    ws = wb["口径说明"]
    all_text = "\n".join(
        str(c.value) for row in ws.iter_rows() for c in row if c.value is not None
    )
    for keyword in ["淘宝/天猫（生意参谋）", "抖音电商（抖店罗盘）", "拼多多（商家后台）",
                    "京东（京东商智）", "京东 GMV 含运费", "支付", "下单"]:
        assert keyword in all_text, f"口径说明缺少：{keyword}"


def test_long_product_name_column_width(seeded_report, export_dir):
    """商品名超长（>100 字符）不破坏列宽（上限 45）。"""
    path = generate_excel_report(seeded_report, export_dir)
    wb = openpyxl.load_workbook(path)
    ws = wb["分SKU"]
    assert ws.column_dimensions["B"].width == 45
    assert ws.max_column >= 9


def test_title_and_warning_block(seeded_report, export_dir):
    path = generate_excel_report(seeded_report, export_dir, company="测试公司")
    wb = openpyxl.load_workbook(path)
    ws = wb["日报总览"]
    assert "测试公司 电商销售日报 2026-09-06 周日" in str(ws["A1"].value)
    # 缺失平台警告块（doudian/pinduoduo 无数据）
    texts = [str(c.value) for row in ws.iter_rows() for c in row if c.value]
    assert any("缺失" in t for t in texts)
