"""利润指标测试（2026-09 补充）：货物成本 / 平台扣点 / 履约成本 / 毛利链 / 成本覆盖率。

业务约定：
- 毛利 = 成交额 - 货成本（SKU成本价×件数）- 平台扣点 - 单件履约成本
- 经营利润 = 毛利 - 推广费
- 成本未知（未映射 SKU / 未填成本价 / 无件数）的行 cogs 为 None，
  绝不按 0 虚报；覆盖率 <100% 时 data_completeness 必须给出警告
"""

from __future__ import annotations

from datetime import date, timedelta
from decimal import Decimal

import pytest
import yaml

from backend.app.config import settings
from backend.app.core.aggregator import build_daily_report
from backend.app.models.entities import AdapterConfig, Sku, Upload
from backend.tests.test_aggregator import _insert_fact

D0 = date(2026, 9, 6)
D1 = D0 - timedelta(days=1)


@pytest.fixture()
def profit_dataset(db_session):
    """已知数据集：
    - SKU1 成本价 30：今日淘宝 A1 gmv1000/退款100/net900/10件/推广50；
      昨日 net720/8件（供环比）
    - 抖店 B1 未映射：今日 net300/3件（成本未知 → 覆盖率 75%）
    """
    up = Upload(filename="t.csv", sha256="profit001", platform_key="taobao",
                status="committed")
    db_session.add(up)
    db_session.flush()

    sku = Sku(sku_code="SKU1", name="测试商品一", category="美妆",
              cost_price=Decimal("30"))
    db_session.add(sku)
    db_session.flush()

    _insert_fact(db_session, up.id, D0, "taobao", "A1", gmv="1000", refund="100",
                 paid_qty=10, buyers=8, visitors=200, ad_cost="50", sku_id=sku.id)
    _insert_fact(db_session, up.id, D1, "taobao", "A1", gmv="800", refund="80",
                 paid_qty=8, buyers=6, visitors=160, sku_id=sku.id)
    _insert_fact(db_session, up.id, D0, "doudian", "B1", gmv="300", refund="0",
                 paid_qty=3, buyers=3, visitors=60, shop_code="S2",
                 shop_name="店铺二", category="食品")
    db_session.commit()
    return {"upload_id": up.id, "sku_id": sku.id}


def _override_taobao_fee(db_session, rate: str, per_unit: str) -> None:
    """以用户覆盖版适配器的方式给淘宝配置扣点与单件履约成本。"""
    builtin = (settings.adapters_dir / "taobao.yaml").read_text(encoding="utf-8")
    data = yaml.safe_load(builtin)
    data["metric_semantics"]["platform_fee_rate"] = rate
    data["metric_semantics"]["fulfillment_cost_per_unit"] = per_unit
    db_session.add(
        AdapterConfig(
            platform_key="taobao",
            version=99,
            yaml_text=yaml.safe_dump(data, allow_unicode=True, sort_keys=False),
            is_builtin=True,
            enabled=True,
        )
    )
    db_session.commit()


def test_profit_chain_default(profit_dataset, db_session):
    """默认配置（扣点 0、履约 0）：毛利 = 成交额 - 货成本。"""
    report = build_daily_report(db_session, D0)
    ov = report["overview"]

    # 成交额 1200；货成本仅 SKU1 部分 = 30×10 = 300（B1 未映射不计）
    assert ov["net_amount"]["value"] == 1200.0
    assert ov["gross_profit"]["value"] == pytest.approx(900.0)
    assert ov["gross_margin"]["value"] == pytest.approx(0.75)
    # 经营利润 = 900 - 推广 50 = 850
    assert ov["operating_profit"]["value"] == pytest.approx(850.0)

    # 成本覆盖率 = 900/1200 = 0.75，且必须给出警告
    dc = report["data_completeness"]
    assert dc["cost_coverage"] == pytest.approx(0.75)
    assert dc["warning"] and "成本价" in dc["warning"]

    # 分SKU 行也带毛利：SKU1 毛利 = 900 - 300 = 600
    mapped = report["by_sku"]["mapped"]
    assert len(mapped) == 1
    assert mapped[0]["metrics"]["gross_profit"]["value"] == pytest.approx(600.0)


def test_profit_dod_compare(profit_dataset, db_session):
    """毛利环比同口径：昨日毛利 = 720 - 30×8 = 480，今日 900 → +87.5%。"""
    report = build_daily_report(db_session, D0)
    ov = report["overview"]
    assert ov["gross_profit"]["dod"] == pytest.approx((900 - 480) / 480, abs=1e-3)


def test_profit_with_fee_override(profit_dataset, db_session):
    """淘宝扣点 10% + 履约 2 元/件：扣点 90、履约 20，毛利再减 110。"""
    _override_taobao_fee(db_session, "0.1", "2")
    report = build_daily_report(db_session, D0)
    ov = report["overview"]

    # 毛利 = 1200 - 300 - 90（900×0.1）- 20（10×2）= 790
    assert ov["gross_profit"]["value"] == pytest.approx(790.0)
    # 经营利润 = 790 - 50 = 740
    assert ov["operating_profit"]["value"] == pytest.approx(740.0)


def test_profit_none_without_any_cost(db_session):
    """完全没有成本价：毛利链整体为 None（不虚报），覆盖率 0 + 警告。"""
    up = Upload(filename="t.csv", sha256="profit002", platform_key="taobao",
                status="committed")
    db_session.add(up)
    db_session.flush()
    _insert_fact(db_session, up.id, D0, "taobao", "A1", gmv="1000", refund="100",
                 paid_qty=10, buyers=8, visitors=200)
    db_session.commit()

    report = build_daily_report(db_session, D0)
    ov = report["overview"]
    assert ov["gross_profit"]["value"] is None
    assert ov["gross_margin"]["value"] is None
    assert ov["operating_profit"]["value"] is None
    dc = report["data_completeness"]
    assert dc["cost_coverage"] == 0.0
    assert dc["warning"] and "成本价" in dc["warning"]


def test_summary_profit_line(profit_dataset, db_session):
    """文字摘要：有成本时出现利润行，无成本时整行省略。"""
    from backend.app.exporters.text_summary import generate_text_summary

    report = build_daily_report(db_session, D0)
    text = generate_text_summary(report, company="测试")
    assert "毛利" in text and "经营利润" in text
    assert len(text) <= 800

    # 无成本数据的报告（手工构造 overview 无 gross 值）→ 不出现利润行
    # （覆盖率警告文案本身会提到"毛利"，故断言利润行的特征前缀"毛利 ¥"）
    report2 = build_daily_report(db_session, D0)
    report2["overview"]["gross_profit"]["value"] = None
    text2 = generate_text_summary(report2, company="测试")
    assert "毛利 ¥" not in text2


def test_excel_profit_rows(profit_dataset, db_session, export_dir):
    """Excel 日报总览含毛利/毛利率/经营利润三行，文件可正常生成。"""
    from openpyxl import load_workbook

    from backend.app.exporters.excel_report import generate_excel_report

    report = build_daily_report(db_session, D0)
    path = generate_excel_report(report, export_dir, company="测试")
    wb = load_workbook(path)
    ws = wb["日报总览"]
    labels = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
    for expected in ("毛利", "毛利率", "经营利润"):
        assert expected in labels, f"日报总览缺少 {expected} 行"
